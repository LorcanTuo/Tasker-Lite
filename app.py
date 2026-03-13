from flask import Flask, render_template, request, redirect, url_for, flash, send_file
from flask_login import LoginManager, UserMixin, login_user, logout_user, login_required, current_user
from werkzeug.security import generate_password_hash, check_password_hash
from models import get_db, init_db
from notifications import notify_ticket
from datetime import datetime, date, timedelta
import config
import io
import os
import csv
import base64
import subprocess
import tempfile
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

app = Flask(__name__)
app.secret_key = config.SECRET_KEY

login_manager = LoginManager(app)
login_manager.login_view = 'login'
login_manager.login_message = 'Please log in to access this page.'
login_manager.login_message_category = 'warning'


class User(UserMixin):
    def __init__(self, id, username):
        self.id = id
        self.username = username


@login_manager.user_loader
def load_user(user_id):
    db = get_db()
    row = db.execute("SELECT id, username FROM users WHERE id=?", (user_id,)).fetchone()
    db.close()
    if row:
        return User(row['id'], row['username'])
    return None


@app.before_request
def before_request():
    init_db()


@app.context_processor
def inject_config():
    """Make config values available in every template."""
    return dict(cfg=config)


# ---------- DASHBOARD ----------

@app.route('/')
@login_required
def dashboard():
    db = get_db()
    today = date.today().isoformat()

    walkin_today = db.execute(
        "SELECT COUNT(*) as c FROM walkin_tickets WHERE date = ?", (today,)
    ).fetchone()['c']

    walkin_total = db.execute("SELECT COUNT(*) as c FROM walkin_tickets").fetchone()['c']
    walkin_open = db.execute("SELECT COUNT(*) as c FROM walkin_tickets WHERE status = 'IN PROGRESS'").fetchone()['c']
    escalated_total = db.execute("SELECT COUNT(*) as c FROM escalated_tickets").fetchone()['c']
    escalated_open = db.execute("SELECT COUNT(*) as c FROM escalated_tickets WHERE status = 'IN PROGRESS'").fetchone()['c']

    recent_walkin = db.execute(
        "SELECT * FROM walkin_tickets ORDER BY date DESC, id DESC LIMIT 10"
    ).fetchall()
    recent_escalated = db.execute(
        "SELECT * FROM escalated_tickets ORDER BY date DESC, id DESC LIMIT 5"
    ).fetchall()

    db.close()
    return render_template('dashboard.html',
        walkin_today=walkin_today, walkin_total=walkin_total, walkin_open=walkin_open,
        escalated_total=escalated_total, escalated_open=escalated_open,
        recent_walkin=recent_walkin, recent_escalated=recent_escalated
    )


# ---------- WALK-IN TICKETS ----------

@app.route('/walkin')
@login_required
def walkin_list():
    db = get_db()
    status_filter = request.args.get('status', '')
    category_filter = request.args.get('category', '')
    search = request.args.get('search', '')

    query = "SELECT * FROM walkin_tickets WHERE 1=1"
    params = []
    if status_filter:
        query += " AND status = ?"
        params.append(status_filter)
    if category_filter:
        query += " AND category = ?"
        params.append(category_filter)
    if search:
        query += " AND (description LIKE ? OR reported_by LIKE ? OR additional_info LIKE ?)"
        params.extend([f'%{search}%'] * 3)
    query += " ORDER BY date DESC, id DESC"

    items = db.execute(query, params).fetchall()
    distinct_categories = [r['category'] for r in db.execute(
        "SELECT DISTINCT category FROM walkin_tickets WHERE category IS NOT NULL AND category != '' ORDER BY category"
    ).fetchall()]
    db.close()
    return render_template('walkin_list.html', items=items, categories=distinct_categories,
        statuses=config.STATUSES, status_filter=status_filter, category_filter=category_filter, search=search)


@app.route('/walkin/new', methods=['GET', 'POST'])
@login_required
def walkin_new():
    if request.method == 'POST':
        db = get_db()
        db.execute('''INSERT INTO walkin_tickets
            (date, category, location, description, reported_by, contact_email, additional_info, assigned_to, resolved_date, status)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
            (
                request.form['date'],
                request.form['category'],
                request.form.get('location', ''),
                request.form['description'],
                request.form.get('reported_by', ''),
                request.form.get('contact_email', ''),
                request.form.get('additional_info', ''),
                request.form.get('assigned_to', ''),
                request.form.get('resolved_date', '') or None,
                request.form.get('status', 'IN PROGRESS'),
            ))
        db.commit()
        db.close()
        notify_ticket(
            config.WALKIN_LABEL_SINGULAR, 'created',
            request.form['description'],
            contact_email=request.form.get('contact_email', ''),
            reported_by=request.form.get('reported_by', ''),
            category=request.form.get('category', ''),
        )
        flash(f'{config.WALKIN_LABEL_SINGULAR} logged.', 'success')
        return redirect(url_for('walkin_list'))
    db = get_db()
    distinct_categories = [r['category'] for r in db.execute(
        "SELECT DISTINCT category FROM walkin_tickets WHERE category IS NOT NULL AND category != '' ORDER BY category"
    ).fetchall()]
    raw_assigned = db.execute("SELECT DISTINCT assigned_to FROM walkin_tickets WHERE assigned_to IS NOT NULL AND assigned_to != ''").fetchall()
    assigned_suggestions = sorted(set(n.strip() for row in raw_assigned for n in row['assigned_to'].split(',') if n.strip()))
    db.close()
    return render_template('walkin_form.html', item=None, categories=distinct_categories,
        assigned_suggestions=assigned_suggestions, statuses=config.STATUSES, today=date.today().isoformat())


@app.route('/walkin/<int:item_id>/edit', methods=['GET', 'POST'])
@login_required
def walkin_edit(item_id):
    db = get_db()
    if request.method == 'POST':
        db.execute('''UPDATE walkin_tickets SET
            date=?, category=?, location=?, description=?, reported_by=?,
            contact_email=?, additional_info=?, assigned_to=?, resolved_date=?, status=?, updated_at=datetime('now')
            WHERE id=?''',
            (
                request.form['date'],
                request.form['category'],
                request.form.get('location', ''),
                request.form['description'],
                request.form.get('reported_by', ''),
                request.form.get('contact_email', ''),
                request.form.get('additional_info', ''),
                request.form.get('assigned_to', ''),
                request.form.get('resolved_date', '') or None,
                request.form.get('status', 'IN PROGRESS'),
                item_id,
            ))
        db.commit()
        db.close()
        notify_ticket(
            config.WALKIN_LABEL_SINGULAR, 'updated',
            request.form['description'],
            contact_email=request.form.get('contact_email', ''),
            reported_by=request.form.get('reported_by', ''),
        )
        flash(f'{config.WALKIN_LABEL_SINGULAR} updated.', 'success')
        return redirect(url_for('walkin_list'))
    item = db.execute("SELECT * FROM walkin_tickets WHERE id=?", (item_id,)).fetchone()
    distinct_categories = [r['category'] for r in db.execute(
        "SELECT DISTINCT category FROM walkin_tickets WHERE category IS NOT NULL AND category != '' ORDER BY category"
    ).fetchall()]
    raw_assigned = db.execute("SELECT DISTINCT assigned_to FROM walkin_tickets WHERE assigned_to IS NOT NULL AND assigned_to != ''").fetchall()
    assigned_suggestions = sorted(set(n.strip() for row in raw_assigned for n in row['assigned_to'].split(',') if n.strip()))
    db.close()
    return render_template('walkin_form.html', item=item, categories=distinct_categories,
        assigned_suggestions=assigned_suggestions, statuses=config.STATUSES, today=date.today().isoformat())


# ---------- ESCALATED TICKETS ----------

@app.route('/escalated')
@login_required
def escalated_list():
    db = get_db()
    status_filter = request.args.get('status', '')
    area_filter = request.args.get('area', '')
    search = request.args.get('search', '')

    query = "SELECT * FROM escalated_tickets WHERE 1=1"
    params = []
    if status_filter:
        query += " AND status = ?"
        params.append(status_filter)
    if area_filter:
        query += " AND area = ?"
        params.append(area_filter)
    if search:
        query += " AND (description LIKE ? OR reported_by LIKE ? OR ticket_number LIKE ? OR additional_info LIKE ?)"
        params.extend([f'%{search}%'] * 4)
    query += " ORDER BY date DESC, id DESC"

    items = db.execute(query, params).fetchall()
    distinct_areas = [r['area'] for r in db.execute(
        "SELECT DISTINCT area FROM escalated_tickets WHERE area IS NOT NULL AND area != '' ORDER BY area"
    ).fetchall()]
    db.close()
    return render_template('escalated_list.html', items=items, areas=distinct_areas,
        statuses=config.STATUSES, status_filter=status_filter, area_filter=area_filter, search=search)


@app.route('/escalated/new', methods=['GET', 'POST'])
@login_required
def escalated_new():
    if request.method == 'POST':
        db = get_db()
        db.execute('''INSERT INTO escalated_tickets
            (date, area, building, office, description, reported_by, contact_email, escalated_by, ticket_number, escalation_date, additional_info, resolved_date, status)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
            (
                request.form['date'],
                request.form['area'],
                request.form.get('building', ''),
                request.form.get('office', ''),
                request.form['description'],
                request.form.get('reported_by', ''),
                request.form.get('contact_email', ''),
                request.form.get('escalated_by', ''),
                request.form.get('ticket_number', ''),
                request.form.get('escalation_date', '') or None,
                request.form.get('additional_info', ''),
                request.form.get('resolved_date', '') or None,
                request.form.get('status', 'IN PROGRESS'),
            ))
        db.commit()
        db.close()
        notify_ticket(
            config.ESCALATED_LABEL_SINGULAR, 'created',
            request.form['description'],
            contact_email=request.form.get('contact_email', ''),
            reported_by=request.form.get('reported_by', ''),
            ticket_number=request.form.get('ticket_number', ''),
        )
        flash(f'{config.ESCALATED_LABEL_SINGULAR} logged.', 'success')
        return redirect(url_for('escalated_list'))
    db = get_db()
    distinct_areas = [r['area'] for r in db.execute(
        "SELECT DISTINCT area FROM escalated_tickets WHERE area IS NOT NULL AND area != '' ORDER BY area"
    ).fetchall()]
    db.close()
    return render_template('escalated_form.html', item=None, areas=distinct_areas,
        statuses=config.STATUSES, today=date.today().isoformat())


@app.route('/escalated/<int:item_id>/edit', methods=['GET', 'POST'])
@login_required
def escalated_edit(item_id):
    db = get_db()
    if request.method == 'POST':
        db.execute('''UPDATE escalated_tickets SET
            date=?, area=?, building=?, office=?, description=?, reported_by=?,
            contact_email=?, escalated_by=?, ticket_number=?, escalation_date=?, additional_info=?,
            resolved_date=?, status=?, updated_at=datetime('now')
            WHERE id=?''',
            (
                request.form['date'],
                request.form['area'],
                request.form.get('building', ''),
                request.form.get('office', ''),
                request.form['description'],
                request.form.get('reported_by', ''),
                request.form.get('contact_email', ''),
                request.form.get('escalated_by', ''),
                request.form.get('ticket_number', ''),
                request.form.get('escalation_date', '') or None,
                request.form.get('additional_info', ''),
                request.form.get('resolved_date', '') or None,
                request.form.get('status', 'IN PROGRESS'),
                item_id,
            ))
        db.commit()
        db.close()
        notify_ticket(
            config.ESCALATED_LABEL_SINGULAR, 'updated',
            request.form['description'],
            contact_email=request.form.get('contact_email', ''),
            reported_by=request.form.get('reported_by', ''),
            ticket_number=request.form.get('ticket_number', ''),
        )
        flash(f'{config.ESCALATED_LABEL_SINGULAR} updated.', 'success')
        return redirect(url_for('escalated_list'))
    item = db.execute("SELECT * FROM escalated_tickets WHERE id=?", (item_id,)).fetchone()
    distinct_areas = [r['area'] for r in db.execute(
        "SELECT DISTINCT area FROM escalated_tickets WHERE area IS NOT NULL AND area != '' ORDER BY area"
    ).fetchall()]
    db.close()
    return render_template('escalated_form.html', item=item, areas=distinct_areas,
        statuses=config.STATUSES, today=date.today().isoformat())


# ---------- QUICK RESOLVE ----------

@app.route('/walkin/<int:item_id>/resolve', methods=['POST'])
@login_required
def walkin_resolve(item_id):
    db = get_db()
    ticket = db.execute("SELECT description, contact_email, reported_by FROM walkin_tickets WHERE id=?", (item_id,)).fetchone()
    db.execute("UPDATE walkin_tickets SET status='RESOLVED', resolved_date=?, updated_at=datetime('now') WHERE id=?",
               (date.today().isoformat(), item_id))
    db.commit()
    db.close()
    if ticket:
        notify_ticket(
            config.WALKIN_LABEL_SINGULAR, 'resolved',
            ticket['description'],
            contact_email=ticket['contact_email'] or '',
            reported_by=ticket['reported_by'] or '',
        )
    flash('Marked as resolved.', 'success')
    return redirect(request.referrer or url_for('walkin_list'))


@app.route('/escalated/<int:item_id>/resolve', methods=['POST'])
@login_required
def escalated_resolve(item_id):
    db = get_db()
    ticket = db.execute("SELECT description, contact_email, reported_by, ticket_number FROM escalated_tickets WHERE id=?", (item_id,)).fetchone()
    db.execute("UPDATE escalated_tickets SET status='RESOLVED', resolved_date=?, updated_at=datetime('now') WHERE id=?",
               (date.today().isoformat(), item_id))
    db.commit()
    db.close()
    if ticket:
        notify_ticket(
            config.ESCALATED_LABEL_SINGULAR, 'resolved',
            ticket['description'],
            contact_email=ticket['contact_email'] or '',
            reported_by=ticket['reported_by'] or '',
            ticket_number=ticket['ticket_number'] or '',
        )
    flash('Marked as resolved.', 'success')
    return redirect(request.referrer or url_for('escalated_list'))


# ---------- REPORTS ----------

@app.route('/reports')
@login_required
def reports():
    return render_template('reports.html')


@app.route('/reports/generate')
@login_required
def generate_report():
    start = request.args.get('start', '')
    end = request.args.get('end', '')
    if not start or not end:
        flash('Please select a date range.', 'warning')
        return redirect(url_for('reports'))

    db = get_db()

    walkin_items = db.execute(
        "SELECT * FROM walkin_tickets WHERE date BETWEEN ? AND ? ORDER BY date",
        (start, end)).fetchall()
    walkin_raised = len(walkin_items)
    walkin_resolved = sum(1 for r in walkin_items if r['status'] == 'RESOLVED')
    walkin_in_progress = walkin_raised - walkin_resolved

    escalated_items = db.execute(
        "SELECT * FROM escalated_tickets WHERE date BETWEEN ? AND ? ORDER BY date",
        (start, end)).fetchall()
    escalated_raised = len(escalated_items)
    escalated_resolved = sum(1 for r in escalated_items if r['status'] == 'RESOLVED')
    escalated_in_progress = escalated_raised - escalated_resolved

    category_counts = {}
    for r in walkin_items:
        t = r['category'] or 'Unknown'
        category_counts[t] = category_counts.get(t, 0) + 1

    db.close()
    return render_template('report_view.html',
        start=start, end=end,
        walkin_raised=walkin_raised, walkin_resolved=walkin_resolved, walkin_in_progress=walkin_in_progress,
        escalated_raised=escalated_raised, escalated_resolved=escalated_resolved, escalated_in_progress=escalated_in_progress,
        category_counts=sorted(category_counts.items(), key=lambda x: x[1], reverse=True),
        walkin_items=walkin_items, escalated_items=escalated_items
    )


# ---------- EXPORT TO EXCEL ----------

@app.route('/export')
@login_required
def export_excel():
    start = request.args.get('start', '')
    end = request.args.get('end', '')
    db = get_db()

    wb = openpyxl.Workbook()

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    def style_header(ws, row=1):
        for cell in ws[row]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', wrap_text=True)
            cell.border = thin_border

    # --- Summary sheet ---
    ws_summary = wb.active
    ws_summary.title = "Status Report"
    ws_summary.append([])
    ws_summary.append([None, f"{config.APP_NAME} — Status Report"])
    ws_summary.append([None, f"Report Period: {start} to {end}" if start else "All Data"])
    ws_summary.append([])

    where_walkin = " WHERE date BETWEEN ? AND ?" if start else ""
    where_escalated = " WHERE date BETWEEN ? AND ?" if start else ""
    params = (start, end) if start else ()

    ws_summary.append([None, "Category", "Raised", "Completed", "In Progress"])
    style_header(ws_summary, 5)

    for label, table, where in [
        (config.ESCALATED_LABEL, "escalated_tickets", where_escalated),
        (config.WALKIN_LABEL, "walkin_tickets", where_walkin),
    ]:
        total = db.execute(f"SELECT COUNT(*) as c FROM {table}{where}", params).fetchone()['c']
        resolved = db.execute(f"SELECT COUNT(*) as c FROM {table}{where}{' AND' if where else ' WHERE'} status='RESOLVED'", params).fetchone()['c']
        ws_summary.append([None, label, total, resolved, total - resolved])

    # --- Walk-in Tickets sheet ---
    ws_walkin = wb.create_sheet(config.WALKIN_LABEL)
    ws_walkin.append(["DATE", "CATEGORY", "LOCATION", "DESCRIPTION", "REPORTED BY",
                      "CONTACT EMAIL", "ADDITIONAL INFO", "ASSIGNED TO", "RESOLVED DATE", "STATUS"])
    style_header(ws_walkin)

    rows = db.execute(f"SELECT * FROM walkin_tickets{where_walkin} ORDER BY date", params).fetchall()
    for r in rows:
        ws_walkin.append([r['date'], r['category'], r['location'], r['description'],
                         r['reported_by'], r['contact_email'], r['additional_info'], r['assigned_to'],
                         r['resolved_date'], r['status']])

    # --- Escalated Tickets sheet ---
    ws_esc = wb.create_sheet(config.ESCALATED_LABEL)
    ws_esc.append(["DATE", "AREA", "BUILDING", "OFFICE", "DESCRIPTION", "REPORTED BY",
                    "CONTACT EMAIL",
                    config.ESCALATED_RAISED_BY_LABEL.upper(), config.ESCALATED_NUMBER_LABEL.upper(),
                    config.ESCALATED_DATE_LABEL.upper(), "ADDITIONAL INFO",
                    "RESOLVED DATE", "STATUS"])
    style_header(ws_esc)

    rows = db.execute(f"SELECT * FROM escalated_tickets{where_escalated} ORDER BY date", params).fetchall()
    for r in rows:
        ws_esc.append([r['date'], r['area'], r['building'], r['office'], r['description'],
                        r['reported_by'], r['contact_email'], r['escalated_by'], r['ticket_number'], r['escalation_date'],
                        r['additional_info'], r['resolved_date'], r['status']])

    db.close()

    # Auto-width columns
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            for cell in col:
                if cell.value:
                    max_len = max(max_len, min(len(str(cell.value)), 40))
            ws.column_dimensions[col[0].column_letter].width = max_len + 2

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"IT_Issues_{start}_to_{end}.xlsx" if start else "IT_Issues_All.xlsx"
    return send_file(output, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ---------- IMPORT FROM EXISTING SPREADSHEET ----------

@app.route('/import', methods=['GET', 'POST'])
@login_required
def import_data():
    if request.method == 'POST':
        if 'file' not in request.files:
            flash('No file selected.', 'warning')
            return redirect(url_for('import_data'))

        file = request.files['file']
        if not file.filename.endswith('.xlsx'):
            flash('Please upload an .xlsx file.', 'warning')
            return redirect(url_for('import_data'))

        try:
            wb = openpyxl.load_workbook(file, data_only=True)
            db = get_db()
            counts = {'walkin': 0, 'escalated': 0}

            def to_date_str(val):
                if val is None:
                    return None
                if isinstance(val, datetime):
                    return val.strftime('%Y-%m-%d')
                return str(val).strip() or None

            # Import Escalated Tickets
            esc_sheet = config.ESCALATED_LABEL
            if esc_sheet in wb.sheetnames:
                ws = wb[esc_sheet]
                for row in range(3, ws.max_row + 1):
                    date_val = to_date_str(ws.cell(row=row, column=1).value)
                    if not date_val:
                        continue
                    db.execute('''INSERT INTO escalated_tickets
                        (date, area, building, office, description, reported_by, escalated_by,
                         ticket_number, escalation_date, additional_info, resolved_date, status)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                        (
                            date_val,
                            str(ws.cell(row=row, column=2).value or '').strip(),
                            str(ws.cell(row=row, column=3).value or '').strip(),
                            str(ws.cell(row=row, column=4).value or '').strip(),
                            str(ws.cell(row=row, column=5).value or '').strip(),
                            str(ws.cell(row=row, column=6).value or '').strip(),
                            str(ws.cell(row=row, column=7).value or '').strip(),
                            str(ws.cell(row=row, column=8).value or '').strip(),
                            to_date_str(ws.cell(row=row, column=9).value),
                            str(ws.cell(row=row, column=10).value or '').strip(),
                            to_date_str(ws.cell(row=row, column=11).value),
                            str(ws.cell(row=row, column=12).value or 'IN PROGRESS').strip(),
                        ))
                    counts['escalated'] += 1

            # Import Walk-in Tickets
            walkin_sheet = config.WALKIN_LABEL
            if walkin_sheet in wb.sheetnames:
                ws = wb[walkin_sheet]
                for row in range(3, ws.max_row + 1):
                    date_val = to_date_str(ws.cell(row=row, column=1).value)
                    if not date_val:
                        continue
                    db.execute('''INSERT INTO walkin_tickets
                        (date, category, location, description, reported_by,
                         additional_info, assigned_to, resolved_date, status)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                        (
                            date_val,
                            str(ws.cell(row=row, column=2).value or '').strip(),
                            str(ws.cell(row=row, column=3).value or '').strip(),
                            str(ws.cell(row=row, column=4).value or '').strip(),
                            str(ws.cell(row=row, column=5).value or '').strip(),
                            str(ws.cell(row=row, column=6).value or '').strip(),
                            str(ws.cell(row=row, column=7).value or '').strip(),
                            to_date_str(ws.cell(row=row, column=8).value),
                            str(ws.cell(row=row, column=9).value or 'IN PROGRESS').strip(),
                        ))
                    counts['walkin'] += 1

            db.commit()
            db.close()
            flash(f"Import complete: {counts['walkin']} walk-in tickets, {counts['escalated']} escalated tickets.", 'success')
        except Exception as e:
            flash(f'Import error: {e}', 'danger')

        return redirect(url_for('import_data'))

    return render_template('import.html')


# ---------- R REPORTS ----------

@app.route('/reports/r')
@login_required
def r_report():
    start = request.args.get('start', '')
    end   = request.args.get('end', '')
    if not start or not end:
        flash('Please select a date range.', 'warning')
        return redirect(url_for('reports'))

    try:
        subprocess.run(['Rscript', '--version'], capture_output=True, check=True, timeout=10)
    except (FileNotFoundError, subprocess.CalledProcessError, subprocess.TimeoutExpired):
        flash('R is not installed or Rscript is not in PATH. Install R from https://cran.r-project.org/ and re-run.', 'danger')
        return redirect(url_for('reports'))

    db = get_db()
    walkin_rows = db.execute(
        "SELECT date, category, location, description, status FROM walkin_tickets "
        "WHERE date BETWEEN ? AND ? ORDER BY date", (start, end)).fetchall()
    escalated_rows = db.execute(
        "SELECT date, area, building, office, description, status FROM escalated_tickets "
        "WHERE date BETWEEN ? AND ? ORDER BY date", (start, end)).fetchall()
    db.close()

    def write_csv(path, rows):
        with open(path, 'w', newline='', encoding='utf-8') as f:
            if rows:
                w = csv.writer(f)
                w.writerow(rows[0].keys())
                w.writerows(rows)

    r_script = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'report.R')

    with tempfile.TemporaryDirectory() as tmpdir:
        walkin_csv = os.path.join(tmpdir, 'walkin.csv')
        escalated_csv = os.path.join(tmpdir, 'escalated.csv')
        write_csv(walkin_csv, walkin_rows)
        write_csv(escalated_csv, escalated_rows)

        result = subprocess.run(
            ['Rscript', r_script, walkin_csv, escalated_csv, tmpdir, start, end],
            capture_output=True, text=True, timeout=180
        )

        if result.returncode != 0:
            flash(f'R report generation failed: {result.stderr[:500]}', 'danger')
            return redirect(url_for('reports'))

        charts = {}
        for name in ['chart_categories', 'chart_timeline', 'chart_status', 'chart_escalated_area']:
            path = os.path.join(tmpdir, f'{name}.png')
            if os.path.exists(path):
                with open(path, 'rb') as f:
                    charts[name] = base64.b64encode(f.read()).decode('utf-8')

    return render_template('r_report.html',
        start=start, end=end, charts=charts,
        walkin_count=len(walkin_rows),
        escalated_count=len(escalated_rows)
    )


# ---------- AUTH ----------

@app.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        db = get_db()
        row = db.execute("SELECT * FROM users WHERE username=?", (username,)).fetchone()
        db.close()
        if row and check_password_hash(row['password_hash'], password):
            login_user(User(row['id'], row['username']), remember=bool(request.form.get('remember')))
            next_page = request.args.get('next')
            return redirect(next_page or url_for('dashboard'))
        flash('Invalid username or password.', 'danger')
    return render_template('login.html')


@app.route('/logout')
@login_required
def logout():
    logout_user()
    flash('You have been logged out.', 'info')
    return redirect(url_for('login'))


# ---------- USER MANAGEMENT ----------

@app.route('/users')
@login_required
def users_list():
    db = get_db()
    users = db.execute("SELECT id, username, created_at FROM users ORDER BY username").fetchall()
    db.close()
    return render_template('users.html', users=users)


@app.route('/users/new', methods=['POST'])
@login_required
def user_new():
    username = request.form.get('username', '').strip()
    password = request.form.get('password', '')
    if not username or not password:
        flash('Username and password are required.', 'warning')
        return redirect(url_for('users_list'))
    db = get_db()
    existing = db.execute("SELECT id FROM users WHERE username=?", (username,)).fetchone()
    if existing:
        db.close()
        flash('Username already exists.', 'danger')
        return redirect(url_for('users_list'))
    db.execute("INSERT INTO users (username, password_hash) VALUES (?, ?)",
               (username, generate_password_hash(password)))
    db.commit()
    db.close()
    flash(f'User "{username}" created.', 'success')
    return redirect(url_for('users_list'))


@app.route('/users/<int:user_id>/delete', methods=['POST'])
@login_required
def user_delete(user_id):
    if user_id == current_user.id:
        flash('You cannot delete your own account.', 'danger')
        return redirect(url_for('users_list'))
    db = get_db()
    db.execute("DELETE FROM users WHERE id=?", (user_id,))
    db.commit()
    db.close()
    flash('User deleted.', 'success')
    return redirect(url_for('users_list'))


@app.route('/change_password', methods=['GET', 'POST'])
@login_required
def change_password():
    if request.method == 'POST':
        current_pw = request.form.get('current_password', '')
        new_pw = request.form.get('new_password', '')
        confirm_pw = request.form.get('confirm_password', '')
        db = get_db()
        row = db.execute("SELECT password_hash FROM users WHERE id=?", (current_user.id,)).fetchone()
        if not check_password_hash(row['password_hash'], current_pw):
            db.close()
            flash('Current password is incorrect.', 'danger')
            return redirect(url_for('change_password'))
        if new_pw != confirm_pw:
            db.close()
            flash('New passwords do not match.', 'danger')
            return redirect(url_for('change_password'))
        if len(new_pw) < 6:
            db.close()
            flash('Password must be at least 6 characters.', 'danger')
            return redirect(url_for('change_password'))
        db.execute("UPDATE users SET password_hash=? WHERE id=?",
                   (generate_password_hash(new_pw), current_user.id))
        db.commit()
        db.close()
        flash('Password changed successfully.', 'success')
        return redirect(url_for('dashboard'))
    return render_template('change_password.html')


if __name__ == '__main__':
    init_db()
    app.run(debug=config.DEBUG, host=config.HOST, port=config.PORT)
